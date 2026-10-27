"""
MediaCompile — Multi-Format Text Extraction, Normalization, and Chunking Engine.
Source of Truth: MediaCompile-prompt.md (Section 1: Supported Content Types & Section 2: Pipeline)

Supported Formats:
- PDF (pypdf): Preserves page numbers for source citations.
- Word DOCX (python-docx): Preserves headings, paragraphs, and tables.
- PowerPoint PPTX (python-pptx): Preserves slide numbers and slide titles.
- Excel XLSX (openpyxl): Preserves sheet names, headers, and row relationships.
- CSV (csv): Preserves column headers and row relationships.
- Plain Text & Markdown (TXT, MD): Encoding-safe extraction preserving headers/bullets.
- Web HTML (BeautifulSoup): Strips boilerplate chrome/ads, captures title, description, canonical URL.
"""

import io
import csv
import re
import logging
from typing import Tuple, List, Dict, Any, Optional
from bs4 import BeautifulSoup
from pypdf import PdfReader
import docx
import pptx
import openpyxl

logger = logging.getLogger(__name__)


def extract_text_from_html(html_bytes: bytes) -> Tuple[str, str, Dict[str, Any]]:
    """
    Extract title, canonical URL, description, and normalized text from HTML.
    Removes boilerplate (scripts, styles, navigation, headers, footers, ads, cookie notices).
    Returns (title, normalized_text, metadata).
    """
    soup = BeautifulSoup(html_bytes, 'html.parser')

    # 1. Capture metadata before stripping tags
    title = ''
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
    elif soup.find('h1'):
        title = soup.find('h1').get_text().strip()

    description = ''
    meta_desc = soup.find('meta', attrs={'name': re.compile(r'description', re.I)}) or \
                soup.find('meta', attrs={'property': 'og:description'})
    if meta_desc and meta_desc.get('content'):
        description = meta_desc['content'].strip()

    canonical_url = ''
    link_canon = soup.find('link', attrs={'rel': 'canonical'}) or \
                 soup.find('meta', attrs={'property': 'og:url'})
    if link_canon:
        canonical_url = link_canon.get('href') or link_canon.get('content', '')

    # 2. Decompose non-content layout boilerplate
    for tag in soup([
        'script', 'style', 'nav', 'header', 'footer', 'aside',
        'noscript', 'iframe', 'svg', 'canvas', 'form'
    ]):
        tag.decompose()

    # Remove common cookie / ad banner classes & ids
    for element in soup.find_all(attrs={'class': re.compile(r'cookie|banner|advert|popup|modal|consent', re.I)}):
        element.decompose()
    for element in soup.find_all(attrs={'id': re.compile(r'cookie|banner|advert|popup|modal|consent', re.I)}):
        element.decompose()

    # 3. Extract main article or body
    body = (
        soup.find('article') or
        soup.find('main') or
        soup.find(attrs={'role': 'main'}) or
        soup.find(attrs={'id': re.compile(r'content|main|article', re.I)}) or
        soup.body or
        soup
    )
    raw_text = body.get_text(separator='\n')

    # Normalize whitespace: collapse consecutive blank lines
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    normalized_text = '\n\n'.join(lines)

    metadata = {
        'description': description[:512] if description else '',
        'canonical_url': canonical_url[:2048] if canonical_url else '',
    }

    return title, normalized_text, metadata


def extract_text_from_pdf(pdf_bytes: bytes) -> Tuple[str, str, List[Dict[str, Any]]]:
    """
    Extract title, text, and per-page content from PDF document using pypdf.
    Preserves 1-indexed page numbers for chatbot source citations.
    Returns (title, normalized_text, pages_data).
    """
    reader = PdfReader(io.BytesIO(pdf_bytes))
    title = ''

    # Read PDF metadata for title if present
    if reader.metadata and reader.metadata.title:
        title = reader.metadata.title.strip()

    pages_data = []
    all_page_texts = []

    for idx, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ''
        except Exception as e:
            logger.warning(f"Failed to extract text from PDF page {idx}: {e}")
            txt = ''

        txt_clean = '\n'.join([line.strip() for line in txt.splitlines() if line.strip()])
        if txt_clean:
            pages_data.append({
                'page_number': idx,
                'text': txt_clean
            })
            all_page_texts.append(f"[Page {idx}]\n{txt_clean}")

    normalized_text = '\n\n'.join(all_page_texts)
    if not title and pages_data:
        # Use first non-empty line of first page text as title fallback
        for line in pages_data[0]['text'].splitlines():
            clean_line = line.strip()
            if clean_line and len(clean_line) > 2:
                title = clean_line[:120]
                break

    return title, normalized_text, pages_data


def extract_text_from_docx(docx_bytes: bytes) -> Tuple[str, str, List[Dict[str, Any]]]:
    """
    Extract title, headings, paragraphs, and tables from Word (.docx) document.
    Returns (title, normalized_text, sections_data).
    """
    doc = docx.Document(io.BytesIO(docx_bytes))
    title = ''

    # Check document properties
    if doc.core_properties and doc.core_properties.title:
        title = doc.core_properties.title.strip()

    sections_data = []
    all_lines = []
    current_heading = ''

    # 1. Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue

        if p.style and p.style.name and p.style.name.startswith('Heading'):
            current_heading = txt
            if not title:
                title = txt[:100]
            all_lines.append(f"## {txt}")
        else:
            all_lines.append(txt)

        sections_data.append({
            'section_title': current_heading,
            'text': txt
        })

    # 2. Extract tables
    for table_idx, table in enumerate(doc.tables, start=1):
        table_rows = []
        headers = []
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            if not any(cells):
                continue
            if row_idx == 0:
                headers = cells
                table_rows.append(" | ".join(headers))
            else:
                if headers and len(headers) == len(cells):
                    row_str = " | ".join(f"{h}: {c}" for h, c in zip(headers, cells) if c)
                else:
                    row_str = " | ".join(cells)
                table_rows.append(f"[Table {table_idx} | Row {row_idx + 1}] {row_str}")

        if table_rows:
            tbl_text = "\n".join(table_rows)
            all_lines.append(tbl_text)
            sections_data.append({
                'section_title': f"Table {table_idx}",
                'text': tbl_text
            })

    normalized_text = '\n\n'.join(all_lines)
    return title, normalized_text, sections_data


def extract_text_from_pptx(pptx_bytes: bytes) -> Tuple[str, str, List[Dict[str, Any]]]:
    """
    Extract title, slide numbers, and slide contents from PowerPoint (.pptx).
    Preserves 1-indexed slide numbers and slide titles.
    Returns (title, normalized_text, slides_data).
    """
    prs = pptx.Presentation(io.BytesIO(pptx_bytes))
    title = ''
    slides_data = []
    all_slide_texts = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = ''
        slide_lines = []

        # Try to get slide title from shapes
        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()
            if not title and idx == 1:
                title = slide_title

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame:
                txt = shape.text_frame.text.strip()
                if txt and txt != slide_title:
                    slide_lines.append(txt)

        # Check for speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_lines.append(f"[Notes: {notes}]")

        header = f"[Slide {idx}" + (f": {slide_title}]" if slide_title else "]")
        body = "\n".join(slide_lines)
        full_slide_str = f"{header}\n{body}".strip()

        slides_data.append({
            'slide_number': idx,
            'slide_title': slide_title,
            'text': full_slide_str
        })
        all_slide_texts.append(full_slide_str)

    normalized_text = '\n\n'.join(all_slide_texts)
    return title, normalized_text, slides_data


def extract_text_from_xlsx(xlsx_bytes: bytes) -> Tuple[str, str, List[Dict[str, Any]]]:
    """
    Extract spreadsheets preserving sheet names, column headers, and row relationships.
    Formats as: [Sheet: <name> | Row <N>] Header1: Value1 | Header2: Value2
    Returns (title, normalized_text, rows_data).
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    rows_data = []
    all_lines = []
    title = ''

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        headers = []
        sheet_lines = [f"[Sheet: {sheet_name}]"]

        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(v).strip() if v is not None else '' for v in row]
            if not any(values):
                continue

            if not headers:
                # First non-empty row serves as column headers
                headers = values
                sheet_lines.append("Headers: " + " | ".join(h for h in headers if h))
                continue

            # Format row with column headers
            row_parts = []
            for h, v in zip(headers, values):
                if v:
                    header_label = h if h else "Column"
                    row_parts.append(f"{header_label}: {v}")

            if row_parts:
                row_str = f"[Sheet: {sheet_name} | Row {row_idx}] " + " | ".join(row_parts)
                sheet_lines.append(row_str)
                rows_data.append({
                    'sheet_name': sheet_name,
                    'row_number': row_idx,
                    'text': row_str
                })

        if len(sheet_lines) > 1:
            all_lines.append("\n".join(sheet_lines))

    normalized_text = '\n\n'.join(all_lines)
    if wb.sheetnames:
        title = f"Spreadsheet: {wb.sheetnames[0]}"

    return title, normalized_text, rows_data


def extract_text_from_csv(csv_bytes: bytes) -> Tuple[str, str, List[Dict[str, Any]]]:
    """
    Extract CSV preserving column headers and row relationships.
    Formats as: [Row <N>] Header1: Value1 | Header2: Value2
    Returns (title, normalized_text, rows_data).
    """
    # Detect encoding
    text = None
    for enc in ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']:
        try:
            text = csv_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = csv_bytes.decode('utf-8', errors='replace')

    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return 'Empty CSV', '', []

    reader = csv.reader(io.StringIO(text))
    headers = []
    rows_data = []
    formatted_lines = []

    for row_idx, row in enumerate(reader, start=1):
        cells = [c.strip() for c in row]
        if not any(cells):
            continue

        if not headers:
            headers = cells
            formatted_lines.append("Headers: " + " | ".join(headers))
            continue

        row_parts = []
        for h, c in zip(headers, cells):
            if c:
                header_label = h if h else "Column"
                row_parts.append(f"{header_label}: {c}")

        if row_parts:
            row_str = f"[Row {row_idx}] " + " | ".join(row_parts)
            formatted_lines.append(row_str)
            rows_data.append({
                'row_number': row_idx,
                'text': row_str
            })

    title = f"CSV Dataset ({len(rows_data)} rows)"
    normalized_text = '\n'.join(formatted_lines)
    return title, normalized_text, rows_data


def extract_text_from_txt(txt_bytes: bytes) -> Tuple[str, str]:
    """
    Extract text from plain text or Markdown document with encoding fallback.
    Returns (title, normalized_text).
    """
    for encoding in ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']:
        try:
            text = txt_bytes.decode(encoding)
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            title = lines[0][:80] if lines else 'Text Note'
            # If line is markdown header, strip #
            title = re.sub(r'^#+\s*', '', title).strip()
            return title, '\n\n'.join(lines)
        except UnicodeDecodeError:
            continue

    text = txt_bytes.decode('utf-8', errors='replace')
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    title = lines[0][:80] if lines else 'Text Document'
    title = re.sub(r'^#+\s*', '', title).strip()
    return title, '\n\n'.join(lines)


def chunk_text(
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 150,
    metadata_extractor: Optional[Any] = None
) -> List[Dict[str, Any]]:
    """
    Split a document string into sliding window text chunks.
    Detects structural headers ([Page N], [Slide N], [Sheet: X | Row N]) to attach metadata.
    Returns list of chunk dicts:
    - chunk_index
    - text
    - start_offset
    - end_offset
    - char_length
    - metadata (page_number, slide_number, slide_title, sheet_name, row_number)
    """
    if not text or not text.strip():
        return []

    text = text.strip()
    total_len = len(text)
    chunks = []
    chunk_index = 0
    start = 0

    while start < total_len:
        end = min(start + chunk_size, total_len)

        # Break cleanly on paragraph or sentence boundary
        if end < total_len:
            p_break = text.rfind('\n\n', start + chunk_size - 120, end)
            if p_break != -1 and p_break > start + 100:
                end = p_break + 2
            else:
                s_break = max(
                    text.rfind('. ', start + chunk_size - 80, end),
                    text.rfind('! ', start + chunk_size - 80, end),
                    text.rfind('? ', start + chunk_size - 80, end)
                )
                if s_break != -1 and s_break > start + 100:
                    end = s_break + 2

        chunk_str = text[start:end].strip()
        if chunk_str:
            meta = {'char_length': len(chunk_str)}

            # Look backwards from start or inside chunk for page/slide/sheet markers
            search_window = text[max(0, start - 200):end]
            
            # Detect Page number
            page_match = list(re.finditer(r'\[Page\s+(\d+)\]', search_window))
            if page_match:
                meta['page_number'] = int(page_match[-1].group(1))

            # Detect Slide number and title
            slide_match = list(re.finditer(r'\[Slide\s+(\d+)(?::\s*([^\]]+))?\]', search_window))
            if slide_match:
                last_slide = slide_match[-1]
                meta['slide_number'] = int(last_slide.group(1))
                if last_slide.group(2):
                    meta['slide_title'] = last_slide.group(2).strip()

            # Detect Sheet name and row
            sheet_match = list(re.finditer(r'\[Sheet:\s*([^\|\]]+)(?:\s*\|\s*Row\s*(\d+))?\]', search_window))
            if sheet_match:
                last_sheet = sheet_match[-1]
                meta['sheet_name'] = last_sheet.group(1).strip()
                if last_sheet.group(2):
                    meta['row_number'] = int(last_sheet.group(2))

            # Detect standalone CSV row number
            if 'row_number' not in meta:
                row_match = list(re.finditer(r'\[Row\s+(\d+)\]', search_window))
                if row_match:
                    meta['row_number'] = int(row_match[-1].group(1))

            # Detect Heading or Section
            heading_match = list(re.finditer(r'(?:^|\n)#{1,6}\s+([^\n]+)', search_window))
            if heading_match:
                section_name = heading_match[-1].group(1).strip()
                meta['heading'] = section_name
                meta['section'] = section_name
            else:
                table_match = list(re.finditer(r'\[Table\s+(\d+)\]', search_window))
                if table_match:
                    meta['section'] = f"Table {table_match[-1].group(1)}"

            chunks.append({
                'chunk_index': chunk_index,
                'text': chunk_str,
                'start_offset': start,
                'end_offset': end,
                'char_length': len(chunk_str),
                'metadata': meta
            })
            chunk_index += 1

        if end >= total_len:
            break

        start = max(start + 1, end - chunk_overlap)

    return chunks
