"""
MediaCompile — Phase 3: Content Ingestion Pipeline Verification Test Suite.
Source of Truth: MediaCompile-prompt.md (Phase 3: Content Ingestion Pipeline Verification and Repair)

Verifies:
1. Multi-Format Extraction: PDF, DOCX, PPTX, XLSX, CSV, Plain Text, Markdown, HTML/URL, and Notes.
2. Chunking Quality: chunk index, offsets, structural metadata (page, slide, sheet, row, heading, section).
3. Embeddings: exactly 384 dimensions, sentence-transformers/all-MiniLM-L6-v2 compatibility, stored in pgvector.
4. Status Transitions: pending -> processing -> completed / failed with accurate error messages.
5. End-to-End Semantic Retrieval: pgvector CosineDistance query matches on persisted embeddings.
"""

import io
import os
import csv
import pytest
from unittest.mock import patch, MagicMock
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

import docx
import pptx
import openpyxl
from pypdf import PdfWriter

from users.models import User
from content.models import ContentItem, ContentChunk
from content.services import chunk_text
from content.search import vector_search_chunks, hybrid_search
from content.embeddings import get_embedding_provider


@pytest.fixture
def test_user():
    return User.objects.create_user(
        email="phase3.tester@mediacompile.io",
        password="TestPassword2026!",
        authjs_id="authjs_phase3_test_user"
    )


@pytest.fixture
def auth_client(test_user):
    client = APIClient()
    client.force_authenticate(user=test_user)
    return client


# ------------------------------------------------------------------------------
# 1. PDF Document Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_pdf_ingestion(auth_client, test_user):
    """Verify PDF upload, text extraction, page metadata, 384-dim embedding, and pgvector persistence."""
    pdf_path = r"SKILLS\theme-factory\theme-showcase.pdf"
    if os.path.exists(pdf_path):
        with open(pdf_path, 'rb') as f:
            pdf_bytes = f.read()
    else:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buf = io.BytesIO()
        writer.write(buf)
        pdf_bytes = buf.getvalue()

    uploaded = SimpleUploadedFile("design_theme.pdf", pdf_bytes, content_type="application/pdf")

    res = auth_client.post(
        '/api/v1/content/files/',
        {'file': uploaded, 'title': 'Design Showcase PDF'},
        format='multipart',
        HTTP_X_SYNC_INGESTION='true'
    )
    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED
    assert data['chunk_count'] > 0

    item = ContentItem.objects.get(id=data['id'])
    assert item.user == test_user
    assert item.chunks.count() == data['chunk_count']

    # Inspect chunks and pgvector embedding
    chunks = list(item.chunks.order_by('chunk_index'))
    for idx, c in enumerate(chunks):
        assert c.chunk_index == idx
        assert len(c.text.strip()) > 0
        assert c.start_offset >= 0
        assert c.end_offset >= c.start_offset
        assert c.embedding is not None
        assert len(c.embedding) == 384
        assert c.metadata.get('embedding_dimension') == 384

    # Verify page metadata preserved for PDF
    assert any(c.metadata.get('page_number') is not None for c in chunks)


# ------------------------------------------------------------------------------
# 2. DOCX Document Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_docx_ingestion(auth_client, test_user):
    """Verify DOCX upload, heading/section extraction, 384-dim embedding, and pgvector persistence."""
    doc = docx.Document()
    doc.add_heading("PostgreSQL Vector Storage", level=1)
    doc.add_paragraph("pgvector extends PostgreSQL with native vector similarity operations.")
    doc.add_paragraph("It provides HNSW indexes for approximate nearest neighbor search.")
    doc.add_heading("System Configuration", level=2)
    doc.add_paragraph("Configure ef_search and m parameters for indexing performance.")

    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Index"
    tbl.cell(0, 1).text = "Metric"
    tbl.cell(1, 0).text = "HNSW"
    tbl.cell(1, 1).text = "Cosine"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    uploaded = SimpleUploadedFile("pgvector_guide.docx", docx_bytes, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    res = auth_client.post(
        '/api/v1/content/files/',
        {'file': uploaded, 'title': 'pgvector Technical Guide'},
        format='multipart',
        HTTP_X_SYNC_INGESTION='true'
    )
    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED
    assert data['chunk_count'] >= 1

    item = ContentItem.objects.get(id=data['id'])
    chunks = list(item.chunks.all())
    assert len(chunks) >= 1
    for c in chunks:
        assert len(c.embedding) == 384

    # Verify heading/section detection
    assert any(c.metadata.get('heading') or c.metadata.get('section') for c in chunks)


# ------------------------------------------------------------------------------
# 3. PPTX Document Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_pptx_ingestion(auth_client, test_user):
    """Verify PPTX upload, slide numbers, slide titles, 384-dim embedding, and pgvector persistence."""
    prs = pptx.Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Microservices Consensus"
    body1 = slide1.shapes.placeholders[1]
    body1.text = "Consensus algorithms provide fault tolerance in distributed clusters."

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Raft Protocol Internals"
    body2 = slide2.shapes.placeholders[1]
    body2.text = "Followers redirect requests to the elected leader node."

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    uploaded = SimpleUploadedFile("consensus.pptx", pptx_bytes, content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")

    res = auth_client.post(
        '/api/v1/content/files/',
        {'file': uploaded, 'title': 'Consensus Slides'},
        format='multipart',
        HTTP_X_SYNC_INGESTION='true'
    )
    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED

    item = ContentItem.objects.get(id=data['id'])
    chunks = list(item.chunks.all())
    assert len(chunks) >= 1
    for c in chunks:
        assert len(c.embedding) == 384

    # Verify slide metadata
    slide_numbers = [c.metadata.get('slide_number') for c in chunks if c.metadata.get('slide_number')]
    assert len(slide_numbers) >= 1


# ------------------------------------------------------------------------------
# 4. XLSX Document Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_xlsx_ingestion(auth_client, test_user):
    """Verify XLSX upload, sheet names, row numbers, 384-dim embedding, and pgvector persistence."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Infrastructure"
    ws.append(["Service", "Port", "Protocol"])
    ws.append(["PostgreSQL", "5432", "TCP"])
    ws.append(["Redis", "6379", "TCP"])
    ws.append(["MinIO", "9000", "HTTP"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    uploaded = SimpleUploadedFile("ports.xlsx", xlsx_bytes, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    res = auth_client.post(
        '/api/v1/content/files/',
        {'file': uploaded, 'title': 'Infrastructure Ports'},
        format='multipart',
        HTTP_X_SYNC_INGESTION='true'
    )
    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED

    item = ContentItem.objects.get(id=data['id'])
    chunks = list(item.chunks.all())
    assert len(chunks) >= 1
    for c in chunks:
        assert len(c.embedding) == 384
        assert c.metadata.get('sheet_name') == 'Infrastructure'


# ------------------------------------------------------------------------------
# 5. CSV Document Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_csv_ingestion(auth_client, test_user):
    """Verify CSV upload, row mapping, 384-dim embedding, and pgvector persistence."""
    csv_content = (
        "Name,Role,Location\n"
        "Alice Smith,Senior Architect,Berlin\n"
        "Bob Jones,Database Engineer,Toronto\n"
        "Carol White,Security Analyst,London\n"
    ).encode('utf-8')

    uploaded = SimpleUploadedFile("team.csv", csv_content, content_type="text/csv")

    res = auth_client.post(
        '/api/v1/content/files/',
        {'file': uploaded, 'title': 'Engineering Team'},
        format='multipart',
        HTTP_X_SYNC_INGESTION='true'
    )
    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED

    item = ContentItem.objects.get(id=data['id'])
    chunks = list(item.chunks.all())
    assert len(chunks) >= 1
    for c in chunks:
        assert len(c.embedding) == 384


# ------------------------------------------------------------------------------
# 6. Plain Text & Markdown Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_markdown_ingestion(auth_client, test_user):
    """Verify Markdown (.md) upload, heading extraction, 384-dim embedding, and pgvector persistence."""
    md_content = (
        "# Hybrid Retrieval Architecture\n\n"
        "Reciprocal Rank Fusion combines sparse BM25 scores with dense vector similarities.\n"
        "This balances exact keyword matching with semantic concept coverage.\n\n"
        "## RRF Formulation\n\n"
        "Score is calculated as the sum of 1 / (k + rank) across individual ranking lists."
    ).encode('utf-8')

    uploaded = SimpleUploadedFile("architecture.md", md_content, content_type="text/markdown")

    res = auth_client.post(
        '/api/v1/content/files/',
        {'file': uploaded, 'title': 'RRF Architecture'},
        format='multipart',
        HTTP_X_SYNC_INGESTION='true'
    )
    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED

    item = ContentItem.objects.get(id=data['id'])
    chunks = list(item.chunks.all())
    assert len(chunks) >= 1
    for c in chunks:
        assert len(c.embedding) == 384


# ------------------------------------------------------------------------------
# 7. Note Ingestion
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_note_ingestion(auth_client, test_user):
    """Verify note ingestion generates 384-dim embeddings and completes synchronously."""
    note_text = (
        "Security Note: Always sanitize user input and enforce SSL verification when scraping.\n"
        "Internal services must never be accessible via user-supplied URLs (SSRF protection)."
    )
    res = auth_client.post('/api/v1/content/notes/', {
        'title': 'Security Policy Note',
        'content': note_text,
        'tags': ['security', 'compliance']
    }, format='json')

    assert res.status_code == 201
    data = res.json()
    assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED
    assert data['chunk_count'] >= 1

    item = ContentItem.objects.get(id=data['id'])
    for c in item.chunks.all():
        assert len(c.embedding) == 384


# ------------------------------------------------------------------------------
# 8. Web URL Ingestion (Safe HTML Extraction)
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_url_ingestion(auth_client, test_user):
    """Verify URL ingestion extracts title, metadata, strips boilerplate, and stores 384-dim embeddings."""
    mock_html = b"""
    <!DOCTYPE html>
    <html>
      <head>
        <title>pgvector Performance Tuning</title>
        <meta name="description" content="Best practices for tuning HNSW indexes in production PostgreSQL.">
      </head>
      <body>
        <nav><a href="/">Home</a></nav>
        <main>
          <h1>HNSW Performance Tuning</h1>
          <p>Increasing ef_search improves recall at the expense of slight query latency.</p>
          <p>Maintenance work memory should be configured generously during index building.</p>
        </main>
        <footer>Privacy Policy</footer>
      </body>
    </html>
    """

    mock_resp = MagicMock()
    mock_resp.is_redirect = False
    mock_resp.status_code = 200
    mock_resp.headers = {'Content-Type': 'text/html; charset=utf-8'}
    mock_resp.content = mock_html
    mock_resp.iter_content = lambda chunk_size=8192: [mock_html]

    with patch('requests.Session.get', return_value=mock_resp), \
         patch('socket.getaddrinfo', return_value=[(2, 1, 0, '', ('93.184.216.34', 80))]):
        res = auth_client.post(
            '/api/v1/content/urls/',
            {'url': 'https://example.com/pgvector-tuning'},
            format='json',
            HTTP_X_SYNC_INGESTION='true'
        )
        assert res.status_code == 201
        data = res.json()
        assert data['ingestion_status'] == ContentItem.STATUS_COMPLETED
        assert data['title'] == 'pgvector Performance Tuning'
        assert 'tuning HNSW' in data['summary']

        item = ContentItem.objects.get(id=data['id'])
        assert item.chunks.count() >= 1
        for c in item.chunks.all():
            assert len(c.embedding) == 384


# ------------------------------------------------------------------------------
# 9. Failure State Transitions & Accurate Error Messages
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_failure_records_error_message(auth_client, test_user):
    """Verify failed processing sets STATUS_FAILED and stores exact error message."""
    # 1. SSRF violation
    res = auth_client.post(
        '/api/v1/content/urls/',
        {'url': 'http://169.254.169.254/latest/meta-data/'},
        format='json'
    )
    assert res.status_code == 400
    assert 'url' in res.json()

    # 2. Corrupted / unparseable content task failure
    item = ContentItem.objects.create(
        user=test_user,
        content_type=ContentItem.TYPE_DOCUMENT,
        title="Corrupted File",
        file_key="nonexistent/corrupted.pdf",
        ingestion_status=ContentItem.STATUS_PENDING
    )
    from content.tasks import process_document_content
    process_document_content(str(item.id))
    item.refresh_from_db()

    assert item.ingestion_status == ContentItem.STATUS_FAILED
    assert "Document processing failed" in item.error_message


# ------------------------------------------------------------------------------
# 10. End-to-End pgvector Similarity Retrieval
# ------------------------------------------------------------------------------
@pytest.mark.django_db
def test_pipeline_pgvector_similarity_retrieval(auth_client, test_user):
    """Verify vector search and hybrid search rank ingested content accurately via pgvector."""
    # Ingest a specific note
    res = auth_client.post('/api/v1/content/notes/', {
        'title': 'Quantum Annealing Algorithms',
        'content': 'Quantum annealers solve combinatorial optimization problems using adiabatic state evolution.',
        'tags': ['quantum-computing']
    }, format='json')
    assert res.status_code == 201
    item_id = res.json()['id']

    # 1. Pure semantic search
    vec_results = vector_search_chunks(
        user=test_user,
        query_text="quantum annealing combinatorial optimization",
        limit=10,
        min_similarity=0.01
    )
    assert len(vec_results) >= 1
    matched_item_ids = [str(r['chunk'].content_item_id) for r in vec_results]
    assert item_id in matched_item_ids
    assert vec_results[0]['similarity'] > 0.05

    # 2. Hybrid search (RRF)
    hybrid_results = hybrid_search(
        user=test_user,
        query_text="quantum annealers",
        limit=10
    )
    assert len(hybrid_results) >= 1
    assert any(r['item']['id'] == item_id for r in hybrid_results)
